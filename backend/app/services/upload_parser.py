"""Parse and validate an uploaded service usage file (CSV or Excel).

Nothing here touches the database. The parser turns the file into rows keyed by
normalized column name, and the validator turns those rows into typed values or a
list of row-level errors. Row numbers match what the user sees in a spreadsheet:
the header is row 1, so the first data row is row 2.
"""

import csv
import io
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook

REQUIRED_COLUMNS = ("week_start", "visits", "students_served", "pounds_distributed")
OPTIONAL_COLUMNS = ("office", "service")
TEMPLATE_COLUMNS = ("office", "service", *REQUIRED_COLUMNS)

SUPPORTED_EXTENSIONS = (".csv", ".xlsx")
MAX_DATA_ROWS = 10_000
MAX_INT = 2**31 - 1
MAX_POUNDS = Decimal("99999999.99")  # fits Numeric(10, 2)
DATE_FORMATS = ("%Y-%m-%d", "%m/%d/%Y")


class UploadFileError(Exception):
    """The file as a whole cannot be processed (wrong type, unreadable, no rows)."""


@dataclass
class RowError:
    row: int
    column: str | None
    message: str


@dataclass
class ParsedRow:
    row: int
    office: str | None
    service: str | None
    week_start: date
    visits: int
    students_served: int
    pounds_distributed: Decimal
    extra: dict[str, str] = field(default_factory=dict)


def normalize_header(name: Any) -> str:
    return str(name or "").strip().lower().replace(" ", "_")


def read_rows(filename: str, content: bytes) -> tuple[list[str], list[tuple[int, dict[str, Any]]]]:
    """Return the normalized header and the non-blank data rows with their row numbers."""
    extension = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension not in SUPPORTED_EXTENSIONS:
        raise UploadFileError("Upload a .csv or .xlsx file.")

    raw_rows = _read_csv(content) if extension == ".csv" else _read_xlsx(content)
    if not raw_rows:
        raise UploadFileError("The file is empty.")

    header = [normalize_header(cell) for cell in raw_rows[0]]
    rows = []
    for index, values in enumerate(raw_rows[1:], start=2):
        if all(_is_blank(value) for value in values):
            continue
        rows.append((index, dict(zip(header, values, strict=False))))

    if not rows:
        raise UploadFileError("The file has a header row but no data rows.")
    if len(rows) > MAX_DATA_ROWS:
        raise UploadFileError(f"The file has more than {MAX_DATA_ROWS:,} data rows.")
    return header, rows


def validate_rows(
    header: list[str], rows: list[tuple[int, dict[str, Any]]]
) -> tuple[list[ParsedRow], list[RowError]]:
    missing = [column for column in REQUIRED_COLUMNS if column not in header]
    if missing:
        return [], [RowError(1, column, "Required column is missing.") for column in missing]

    duplicates = sorted({column for column in header if column and header.count(column) > 1})
    if duplicates:
        return [], [RowError(1, column, "Column appears more than once.") for column in duplicates]

    extra_columns = [
        column for column in header if column and column not in TEMPLATE_COLUMNS
    ]
    parsed: list[ParsedRow] = []
    errors: list[RowError] = []

    for row_number, values in rows:
        row_errors: list[RowError] = []
        week_start = _check(row_number, values, "week_start", _parse_date, row_errors)
        visits = _check(row_number, values, "visits", _parse_count, row_errors)
        students_served = _check(row_number, values, "students_served", _parse_count, row_errors)
        pounds = _check(row_number, values, "pounds_distributed", _parse_pounds, row_errors)

        errors.extend(row_errors)
        if row_errors:
            continue

        extra = {
            column: _text(values.get(column))
            for column in extra_columns
            if not _is_blank(values.get(column))
        }
        parsed.append(
            ParsedRow(
                row=row_number,
                office=_text(values.get("office")) or None,
                service=_text(values.get("service")) or None,
                week_start=week_start,
                visits=visits,
                students_served=students_served,
                pounds_distributed=pounds,
                extra=extra,
            )
        )

    # The same office, service and week twice in one file would double count.
    seen: dict[tuple[str | None, str | None, date], int] = {}
    for row in parsed:
        key = (row.office, row.service, row.week_start)
        if key in seen:
            errors.append(
                RowError(
                    row.row,
                    "week_start",
                    f"Duplicate of row {seen[key]} (same office, service, and week).",
                )
            )
        else:
            seen[key] = row.row

    errors.sort(key=lambda error: error.row)
    return parsed, errors


def _check(
    row_number: int,
    values: dict[str, Any],
    column: str,
    parse: Callable[[Any], Any],
    row_errors: list[RowError],
) -> Any:
    try:
        return parse(values.get(column))
    except ValueError as error:
        row_errors.append(RowError(row_number, column, str(error)))
        return None


def _read_csv(content: bytes) -> list[list[Any]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Excel on Windows often saves CSVs in this encoding.
        text = content.decode("cp1252", errors="replace")
    return [row for row in csv.reader(io.StringIO(text))]


def _read_xlsx(content: bytes) -> list[list[Any]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as error:  # openpyxl raises several unrelated types for bad files
        raise UploadFileError("The file could not be read as an Excel workbook.") from error
    try:
        sheet = workbook.worksheets[0]
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _parse_date(value: Any) -> date:
    if _is_blank(value):
        raise ValueError("Value is required.")
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    for date_format in DATE_FORMATS:
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            continue
    raise ValueError(f"'{text}' is not a date. Use YYYY-MM-DD, for example 2025-10-13.")


def _parse_count(value: Any) -> int:
    if _is_blank(value):
        raise ValueError("Value is required.")
    if isinstance(value, bool):
        raise ValueError("Must be a whole number.")
    if isinstance(value, float):
        if not value.is_integer():
            raise ValueError(f"'{value:g}' must be a whole number.")
        number = int(value)
    elif isinstance(value, int):
        number = value
    else:
        text = _text(value).replace(",", "")
        try:
            number = int(text)
        except ValueError:
            try:
                as_decimal = Decimal(text)
            except InvalidOperation:
                raise ValueError(f"'{text}' is not a number.") from None
            if not as_decimal.is_finite() or as_decimal != as_decimal.to_integral_value():
                raise ValueError(f"'{text}' must be a whole number.") from None
            number = int(as_decimal)
    if number < 0:
        raise ValueError("Cannot be negative.")
    if number > MAX_INT:
        raise ValueError("Number is too large.")
    return number


def _parse_pounds(value: Any) -> Decimal:
    if _is_blank(value):
        raise ValueError("Value is required. Use 0 if no food was distributed.")
    if isinstance(value, bool):
        raise ValueError("Must be a number.")
    if isinstance(value, float):
        # Excel formulas can leave float noise such as 120.30000000000001.
        text = repr(round(value, 6))
    else:
        text = _text(value).replace(",", "")
    try:
        number = Decimal(text)
    except InvalidOperation:
        raise ValueError(f"'{_text(value)}' is not a number.") from None
    if not number.is_finite():
        raise ValueError(f"'{_text(value)}' is not a number.")
    if number < 0:
        raise ValueError("Cannot be negative.")
    if number > MAX_POUNDS:
        raise ValueError("Number is too large.")
    if number != number.quantize(Decimal("0.01")):
        raise ValueError("Use at most two decimal places.")
    return number.quantize(Decimal("0.01"))
